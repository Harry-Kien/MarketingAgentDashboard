"""
Báo cáo hiệu suất theo nhân viên + xuất Excel.

MỖI CON SỐ ĐO GÌ — VÀ VÌ SAO ĐO ĐÚNG THẾ
-----------------------------------------
- **Hội thoại được giao**: số hội thoại có lượt giao (`conversation_
  assignments`) cho người này bắt đầu trong khoảng ngày.
- **Tin đã gửi**: tin `messages.gui_boi = người này` (0030). Tin do agent
  soạn rồi người bấm duyệt KHÔNG tính — đó là việc duyệt, không phải việc
  viết, và gộp vào là thưởng cho người chỉ bấm nút.
- **Phản hồi đầu (trung vị, p90)**: từ lúc ĐƯỢC GIAO tới tin đầu tiên
  CHÍNH NGƯỜI ẤY gửi trong hội thoại đó. Không dùng
  `conversations.first_responded_at`: cột ấy tính cả tin agent, nên nhân
  viên nhận hội thoại agent đã trả lời sẽ có "phản hồi 0 giây" không công.
  Trung vị + p90 chứ không trung bình: một ca để quên qua đêm kéo trung bình
  lên vài tiếng và che mất cả ngày làm nhanh.
- **Chưa phản hồi**: lượt giao mà người ấy chưa gửi tin nào sau đó — con số
  quản lý cần nhìn đầu tiên.
- **Đơn từ hội thoại được giao**: đơn (trừ đã huỷ) tạo SAU lúc giao, trong
  hội thoại được giao. Là chỉ báo "chốt đơn", không phải quy công tuyệt đối:
  agent có thể đã tư vấn trước đó.
- **Trễ SLA**: sự kiện `*_breached` trên hội thoại được giao.

Ngày tính theo giờ Việt Nam: "hôm nay" của chủ shop là 0h–24h giờ VN, không
phải giờ UTC (lệch 7 tiếng, cắt đôi ca tối).
"""
from __future__ import annotations

import io
from datetime import date, timedelta
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from agent import db

from .routes import can_quyen

router = APIRouter(prefix="/api/bao-cao", tags=["bao-cao"])

MUI_GIO = "Asia/Ho_Chi_Minh"
TOI_DA_NGAY = 366

_SQL = f"""
WITH khoang AS (
    SELECT ($1::date)::timestamp AT TIME ZONE '{MUI_GIO}' AS tu,
           ($2::date + 1)::timestamp AT TIME ZONE '{MUI_GIO}' AS den
),
giao AS (
    SELECT ca.assigned_user_id AS uid, ca.conversation_id, ca.started_at
    FROM conversation_assignments ca, khoang
    WHERE ca.assigned_user_id IS NOT NULL
      AND ca.started_at >= khoang.tu AND ca.started_at < khoang.den
),
phan_hoi AS (
    SELECT g.uid,
           extract(epoch FROM (
               SELECT min(m.created_at) FROM messages m
               WHERE m.conversation_id = g.conversation_id
                 AND m.gui_boi = g.uid AND m.created_at >= g.started_at
           ) - g.started_at) AS giay
    FROM giao g
),
tin AS (
    SELECT m.gui_boi AS uid, count(*) AS so
    FROM messages m, khoang
    WHERE m.gui_boi IS NOT NULL
      AND m.created_at >= khoang.tu AND m.created_at < khoang.den
    GROUP BY 1
),
don AS (
    SELECT g.uid, count(DISTINCT o.id) AS so
    FROM giao g JOIN orders o
      ON o.conversation_id = g.conversation_id
     AND o.created_at >= g.started_at AND o.trang_thai <> 'da_huy'
    GROUP BY 1
),
tre AS (
    SELECT g.uid, count(*) AS so
    FROM giao g JOIN sla_events s
      ON s.conversation_id = g.conversation_id AND s.kind IN ('first_response_breached', 'resolution_breached')
    GROUP BY 1
),
tong AS (
    SELECT n.id, n.ho_ten, n.ten_dang_nhap,
           (SELECT count(DISTINCT conversation_id) FROM giao WHERE uid = n.id) AS hoi_thoai,
           coalesce(tin.so, 0) AS tin_gui,
           (SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY giay)
              FROM phan_hoi WHERE uid = n.id AND giay IS NOT NULL) AS phan_hoi_trung_vi,
           (SELECT percentile_cont(0.9) WITHIN GROUP (ORDER BY giay)
              FROM phan_hoi WHERE uid = n.id AND giay IS NOT NULL) AS phan_hoi_p90,
           (SELECT count(*) FROM phan_hoi WHERE uid = n.id AND giay IS NULL) AS chua_phan_hoi,
           coalesce(don.so, 0) AS don,
           coalesce(tre.so, 0) AS tre_sla
    FROM nguoi_dung n
    LEFT JOIN tin ON tin.uid = n.id
    LEFT JOIN don ON don.uid = n.id
    LEFT JOIN tre ON tre.uid = n.id
)
SELECT * FROM tong
-- Người bị khoá vẫn hiện nếu CÓ số liệu trong khoảng: nghỉ việc giữa tháng
-- không xoá được công của nửa tháng đầu.
WHERE hoi_thoai > 0 OR tin_gui > 0
   OR id IN (SELECT id FROM nguoi_dung WHERE NOT khoa)
ORDER BY tin_gui DESC, hoi_thoai DESC, ho_ten
"""


def _khoang(tu: date | None, den: date | None) -> tuple[date, date]:
    den = den or date.today()
    tu = tu or den - timedelta(days=6)
    if tu > den:
        raise HTTPException(422, "Ngày bắt đầu phải trước ngày kết thúc.")
    if (den - tu).days > TOI_DA_NGAY:
        raise HTTPException(422, f"Khoảng tối đa {TOI_DA_NGAY} ngày.")
    return tu, den


def _so(x) -> float | None:
    return None if x is None else round(float(x), 1)


async def _lay(tu: date, den: date) -> list[dict[str, Any]]:
    rows = await db.fetch(_SQL, tu, den)
    return [{
        "id": str(r["id"]), "ho_ten": r["ho_ten"] or r["ten_dang_nhap"],
        "ten_dang_nhap": r["ten_dang_nhap"], "hoi_thoai": r["hoi_thoai"],
        "tin_gui": r["tin_gui"], "phan_hoi_trung_vi_giay": _so(r["phan_hoi_trung_vi"]),
        "phan_hoi_p90_giay": _so(r["phan_hoi_p90"]), "chua_phan_hoi": r["chua_phan_hoi"],
        "don": r["don"], "tre_sla": r["tre_sla"],
    } for r in rows]


@router.get("/nhan-vien")
async def bao_cao_nhan_vien(
    tu: date | None = Query(None), den: date | None = Query(None),
    _nguoi: dict = Depends(can_quyen("bao_cao.doc")),
) -> dict[str, Any]:
    tu, den = _khoang(tu, den)
    return {"tu": tu.isoformat(), "den": den.isoformat(), "nhan_vien": await _lay(tu, den)}


COT = [("Nhân viên", "ho_ten"), ("Tên đăng nhập", "ten_dang_nhap"),
       ("Hội thoại được giao", "hoi_thoai"), ("Tin đã gửi", "tin_gui"),
       ("Phản hồi đầu — trung vị (giây)", "phan_hoi_trung_vi_giay"),
       ("Phản hồi đầu — p90 (giây)", "phan_hoi_p90_giay"),
       ("Chưa phản hồi", "chua_phan_hoi"), ("Đơn từ hội thoại được giao", "don"),
       ("Trễ SLA", "tre_sla")]


@router.get("/nhan-vien.xlsx")
async def xuat_excel(
    tu: date | None = Query(None), den: date | None = Query(None),
    nguoi: dict = Depends(can_quyen("bao_cao.doc")),
) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    tu, den = _khoang(tu, den)
    dong = await _lay(tu, den)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hiệu suất nhân viên"
    ws.append([f"Hiệu suất nhân viên {tu.isoformat()} → {den.isoformat()} (giờ Việt Nam)"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append([t for t, _ in COT])
    for o in ws[3]:
        o.font = Font(bold=True)
    for d in dong:
        ws.append([d[k] for _, k in COT])
    for i, (t, _) in enumerate(COT, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = max(14, len(t) + 2)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await db.log_event("bao_cao.xuat_excel", actor=nguoi.get("ten_dang_nhap") or "?",
                       loai="nhan_vien", tu=tu.isoformat(), den=den.isoformat())
    ten = f"hieu-suat-nhan-vien_{tu.isoformat()}_{den.isoformat()}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{ten}"'})
