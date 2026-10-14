"""
Báo cáo hiệu suất nhân viên — dữ liệu dựng tay, con số biết trước.

Một báo cáo hiệu suất sai là một lần đánh giá sai một người thật. Nên mỗi
con số ở đây được dựng từ dữ liệu mà kết quả đúng tính được bằng tay.
"""
from __future__ import annotations

import asyncio
import io
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import UUID, uuid4

import asyncpg
from fastapi import FastAPI
from fastapi.testclient import TestClient

from conftest import nguoi_thu

ROOT = Path(__file__).resolve().parent.parent
LUOC_DO = ROOT / "agent" / "schema.sql"
A = UUID("00000000-0000-0000-0000-00000000000a")
B = UUID("00000000-0000-0000-0000-00000000000b")


def _app(url: str, nguoi: dict) -> FastAPI:
    from agent import db
    from agent.api.bao_cao_nhan_vien import router
    from agent.api.routes import nguoi_da_dang_nhap
    from agent.migrations.runner import apply_all

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[nguoi_da_dang_nhap] = lambda: nguoi

    async def _codec(conn):
        await conn.set_type_codec("jsonb", encoder=json.dumps,
                                  decoder=json.loads, schema="pg_catalog")

    @app.on_event("startup")
    async def _mo():
        db._pool = await asyncpg.create_pool(url, min_size=1, max_size=3, init=_codec)
        async with db._pool.acquire() as conn:
            await conn.execute(LUOC_DO.read_text(encoding="utf-8"))
            await apply_all(conn)

    @app.on_event("shutdown")
    async def _dong():
        if db._pool is not None:
            await db._pool.close()
            db._pool = None

    return app


async def _hoi_thoai(conn, ext):
    tk = await conn.fetchval(
        "INSERT INTO channel_accounts (channel, display_name, status, external_account_id) "
        "VALUES ('webchat',$1,'active',$1) RETURNING id", f"web-{ext}")
    kh = await conn.fetchval("INSERT INTO contacts (display_name) VALUES ($1) RETURNING id", ext)
    d = await conn.fetchval(
        "INSERT INTO contact_points (contact_id, channel_account_id, external_user_id) "
        "VALUES ($1,$2,$3) RETURNING id", kh, tk, ext)
    return await conn.fetchval(
        "INSERT INTO conversations (account_id, channel, external_id, customer_ref, "
        "customer_name, contact_id, contact_point_id) "
        "VALUES ($1,'webchat',$2,$2,$2,$3,$4) RETURNING id", tk, ext, kh, d)


def _dung(url: str) -> None:
    """
    A: được giao h1 (trả lời sau 60s, có 1 đơn) và h2 (chưa trả lời);
       gửi thêm 1 tin ở h1 → 2 tin.
    B: được giao h3, trả lời sau 300s; h3 trễ SLA phản hồi đầu.
    Agent gửi 1 tin ở h1 (gui_boi NULL) — KHÔNG được tính cho ai.
    """
    async def chay():
        conn = await asyncpg.connect(url)
        try:
            for uid, ten in ((A, "a"), (B, "b")):
                await conn.execute(
                    "INSERT INTO nguoi_dung (id, ten_dang_nhap, mat_khau_bam, ho_ten) "
                    "VALUES ($1,$2,'x',$3)", uid, ten, ten.upper())
            t0 = datetime.now(timezone.utc) - timedelta(hours=2)
            h1, h2, h3 = [await _hoi_thoai(conn, x) for x in ("h1", "h2", "h3")]
            for h, u in ((h1, A), (h2, A), (h3, B)):
                await conn.execute(
                    "INSERT INTO conversation_assignments (conversation_id, assigned_user_id, "
                    "source, reason, started_at) VALUES ($1,$2,'manual','kt',$3)", h, u, t0)

            async def tin(h, role, gui_boi, sau):
                await conn.execute(
                    "INSERT INTO messages (conversation_id, role, content, gui_boi, created_at) "
                    "VALUES ($1,$2,'x',$3,$4)", h, role, gui_boi, t0 + timedelta(seconds=sau))
            await tin(h1, "agent", None, 10)          # agent trả lời trước — không tính
            await tin(h1, "staff", A, 60)
            await tin(h1, "staff", A, 600)
            await tin(h3, "staff", B, 300)
            await conn.execute(
                "INSERT INTO orders (ma_don, conversation_id, channel, khach_ten, khach_sdt, "
                "khach_dia_chi, items, tong_tien, trang_thai, created_at) "
                "VALUES ('KT1',$1,'webchat','k','0900000000','x','[]'::jsonb,100000,'da_chot',$2)",
                h1, t0 + timedelta(minutes=20))
            await conn.execute(
                "INSERT INTO sla_events (conversation_id, kind) VALUES ($1,'first_response_breached')", h3)
        finally:
            await conn.close()
    asyncio.run(chay())


def _theo_ten(ds):
    return {d["ten_dang_nhap"]: d for d in ds}


def test_con_so_dung_tung_nguoi(csdl_kiem_thu):
    with TestClient(_app(csdl_kiem_thu, nguoi_thu("bao_cao.doc"))) as k:
        _dung(csdl_kiem_thu)
        r = k.get("/api/bao-cao/nhan-vien")
        assert r.status_code == 200, r.text
        ds = _theo_ten(r.json()["nhan_vien"])
        a, b = ds["a"], ds["b"]
        assert (a["hoi_thoai"], a["tin_gui"], a["chua_phan_hoi"], a["don"], a["tre_sla"]) == (2, 2, 1, 1, 0)
        # 60 giây, KHÔNG phải 10 giây của agent.
        assert a["phan_hoi_trung_vi_giay"] == 60.0
        assert (b["hoi_thoai"], b["tin_gui"], b["chua_phan_hoi"], b["don"], b["tre_sla"]) == (1, 1, 0, 0, 1)
        assert b["phan_hoi_trung_vi_giay"] == 300.0


def test_ngoai_khoang_ngay_thi_khong_tinh(csdl_kiem_thu):
    with TestClient(_app(csdl_kiem_thu, nguoi_thu("bao_cao.doc"))) as k:
        _dung(csdl_kiem_thu)
        ds = _theo_ten(k.get("/api/bao-cao/nhan-vien?tu=2020-01-01&den=2020-01-31")
                       .json()["nhan_vien"])
        assert ds["a"]["tin_gui"] == 0 and ds["a"]["hoi_thoai"] == 0


def test_khoang_sai_la_422(csdl_kiem_thu):
    with TestClient(_app(csdl_kiem_thu, nguoi_thu("bao_cao.doc"))) as k:
        assert k.get("/api/bao-cao/nhan-vien?tu=2026-02-01&den=2026-01-01").status_code == 422
        assert k.get("/api/bao-cao/nhan-vien?tu=2020-01-01&den=2026-01-01").status_code == 422


def test_xuat_excel_khop_json(csdl_kiem_thu):
    from openpyxl import load_workbook

    with TestClient(_app(csdl_kiem_thu, nguoi_thu("bao_cao.doc"))) as k:
        _dung(csdl_kiem_thu)
        js = _theo_ten(k.get("/api/bao-cao/nhan-vien").json()["nhan_vien"])
        r = k.get("/api/bao-cao/nhan-vien.xlsx")
        assert r.status_code == 200
        assert "attachment" in r.headers["content-disposition"]
        ws = load_workbook(io.BytesIO(r.content)).active
        dau = [c.value for c in ws[3]]
        dong = {row[1]: dict(zip(dau, row, strict=True)) for row in ws.iter_rows(min_row=4, values_only=True)}
        assert dong["a"]["Tin đã gửi"] == js["a"]["tin_gui"] == 2
        assert dong["b"]["Trễ SLA"] == 1


def test_thieu_quyen_bao_cao_la_403(csdl_kiem_thu):
    with TestClient(_app(csdl_kiem_thu, nguoi_thu("hoi_thoai.doc"))) as k:
        assert k.get("/api/bao-cao/nhan-vien").status_code == 403
        assert k.get("/api/bao-cao/nhan-vien.xlsx").status_code == 403


def test_tin_nhan_vien_ghi_nguoi_gui(csdl_kiem_thu):
    """Đường gửi thật (OutboundService) phải ghi gui_boi — không thì báo cáo luôn 0."""
    from agent import db
    from agent.omnichannel.outbound_service import OutboundService, PostgresOutboundRepository

    async def chay():
        async def _codec(conn):
            await conn.set_type_codec("jsonb", encoder=json.dumps,
                                      decoder=json.loads, schema="pg_catalog")
        pool = await asyncpg.create_pool(csdl_kiem_thu, min_size=1, max_size=2, init=_codec)
        cu = db._pool
        db._pool = pool
        try:
            async with pool.acquire() as conn:
                await conn.execute(LUOC_DO.read_text(encoding="utf-8"))
                from agent.migrations.runner import apply_all
                await apply_all(conn)
                await conn.execute(
                    "INSERT INTO nguoi_dung (id, ten_dang_nhap, mat_khau_bam, ho_ten) "
                    "VALUES ($1,'a','x','A')", A)
                h = await _hoi_thoai(conn, "g1")
            q = await OutboundService(PostgresOutboundRepository()).queue_text(
                conversation_id=h, role="staff", text="chào chị",
                idempotency_key=f"kt-{uuid4()}", gui_boi=A)
            async with pool.acquire() as conn:
                return await conn.fetchval("SELECT gui_boi FROM messages WHERE id=$1", q.message_id)
        finally:
            db._pool = cu
            await pool.close()
    assert asyncio.run(chay()) == A
